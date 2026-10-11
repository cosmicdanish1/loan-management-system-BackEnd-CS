"""
Builds the Migration Pattern Register (ONE Excel for all migration work) from:
  reports/pattern-scan-before.json  / pattern-scan-after.json   (read-only SQL scan of every open loan account: patterns P01-P12, issues I1-I7)
  reports/*-phase2.log                                           (flags raised by the Phase 2 replay, per member)
  reports/reset-*-members.txt                                    (which members were deleted and migrated again)
Usage: python _build_pattern_register.py [after-scan-label]      (default label: after-all, falls back to after-demand-only)
"""
import json, re, os, sys
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.join(HERE, 'reports')
OUT = os.path.abspath(os.path.join(HERE, '..', '..', '..', 'docs', 'final-documentation', 'loan_slot_consolidation', 'Migration_Pattern_Register.xlsx'))
if os.environ.get('OUT_SUFFIX'): OUT = OUT.replace('.xlsx', os.environ['OUT_SUFFIX'] + '.xlsx')
AFTER_LABEL = sys.argv[1] if len(sys.argv) > 1 else ('after-all' if os.path.exists(os.path.join(R, 'pattern-scan-after-all.json')) else 'after-demand-only')

def load(name):
    p = os.path.join(R, name)
    return json.load(open(p, encoding='utf8')) if os.path.exists(p) else None

def members_file(name):
    p = os.path.join(R, name)
    return {l.strip() for l in open(p) if l.strip()} if os.path.exists(p) else set()

before = load('pattern-scan-before.json')
after = load(f'pattern-scan-{AFTER_LABEL}.json') or []
reset_groups = {
    'Single member re-migration (2026-10-03/04)': {'610032487', '610032274'},
    'Pilot of 30 members': members_file('reset-first-due-pilot-30.txt'),
    'Demand-only pattern members': members_file('reset-demand-only-members.txt'),
    'Other members with no data exception': members_file('reset-others-clean-members.txt'),
}
reset_all = set().union(*reset_groups.values())
group_of = {}
for g, ms in reset_groups.items():
    for m in ms:
        group_of.setdefault(m, g)

PATS = [
 ('P01', 'p01_demand_only', 'Demand only', 'Every repayment in the current loan came through demand-master (D) vouchers; no cash or journal credits.'),
 ('P02', 'p02_cash_only', 'Cash only', 'Repaid only by cash receipts (R vouchers), at least 3, no demand payments.'),
 ('P03', 'p03_demand_and_cash_same_month', 'Demand and cash in the same month', 'In at least 2 months the member paid through the demand master AND in cash.'),
 ('P04', 'p04_topups_while_paying', 'Top-up while still paying', 'A new loan amount was added (payment voucher) while the member was already repaying.'),
 ('P05', 'p05_emi_amount_changes', 'EMI amount keeps changing', 'At least 4 different demand EMI amounts (6+ demand payments).'),
 ('P06', 'p06_cash_lump_sum', 'Cash lump sum', 'A cash receipt of 6 times the usual demand EMI or more.'),
 ('P07', 'p07_silent_12_months', 'Silent 12+ months', 'No payment for 12 months or more (or never) on a loan that is still open.'),
 ('P08', 'p08_journal_entries', 'Journal entries in the loan', 'Journal (J voucher) entries touch this loan account.'),
 ('P09', 'p09_consolidated_loan', 'Consolidated loan / old-EMI watch', 'The loan was consolidated (or has an old-EMI watch window).'),
 ('P10', 'p10_never_paid', 'Never paid since taken', 'No credit at all since the loan was taken (loan older than 60 days).'),
 ('P11', 'p11_gap_in_demand_payments', 'Gap in demand payments', 'More than 45 days between two demand payments.'),
 ('P12', 'p12_cash_receipt_without_interest', 'Cash receipt without interest', 'A cash receipt that carries no interest leg.'),
]
ISS = [
 ('I1', 'i1_first_due_month_late', 'First EMI month stored one month late', 'The stored schedule starts one month after the real first EMI month (migration fault, fixed in the script).'),
 ('I2', 'i2_first_emi_marked_old_loan', "First EMI marked as the old loan's deduction", 'The first real EMI was flagged as predecessor payroll because of I1; it was left out of the principal paid.'),
 ('I3', 'i3_no_schedule_version', 'No schedule version', 'The open loan has no stored schedule (never replayed or not inferable); the older date formula is used.'),
 ('I4', 'i4_replay_failed_for_member', 'Earlier replay failed for the member', 'The earlier Phase 2 replay stopped on a source-data exception for this member.'),
 ('I5', 'i5_record_open_ledger_closed', 'Loan record open, ledger closed', 'The loan record still shows a balance but the ledger says the loan is repaid.'),
 ('I6', 'i6_ledger_record_differ', 'Ledger and loan record differ', 'Ledger balance and loan record balance differ by more than Rs 10.'),
 ('I7', 'i7_ledger_open_record_closed', 'Ledger open, loan record closed', 'The ledger still owes money but the loan record shows the loan closed.'),
]

FLAG_TEXT = {
    'SCHEDULE_NOT_INFERABLE': 'Schedule cannot be inferred: no normal recurring principal payment after the first-EMI month (the EMI has to be confirmed by hand)',
    'MIXED_VOUCHER_INTEREST': 'Regular and emergency loan paid on one voucher: the interest split between them cannot be known',
    'CREDIT_NO_CASE': 'A payment has no open loan case to be applied to (paid more than owed, or the loan list is incomplete)',
    'CREDIT_BEFORE_DISBURSED': 'A payment is dated before the loan it should reduce was disbursed',
    'CONSOLIDATION_UNRESOLVED': 'A consolidation (journal pair) could not be matched to a loan case',
    'OTHER': 'Other data exception, see the Phase 2 log',
    'MANUAL_CORRECTION': 'Known legacy ledger fault (credit posted to the wrong loan type) - corrected by the migration from manual-corrections.json; flagged for manual review',
    'RECONCILED_TO_LEDGER': 'Loan record was set to the ledger closing balance because the legacy loan amount / consolidation did not tie (information; review list)',
    'INFO': 'Information from the replay (partial journal transfer, credit before next loan, ledger-based reopen)',
    'LEDGER_OWES_NO_CASE': 'Ledger still owes but no usable legacy loan row exists (manual)',
    'OVERPAYMENT': 'Overpayment after the loan was fully repaid (information only; shown in the Loan Position panel)',
    'PLACEHOLDER_LOAN': 'Zero-amount placeholder loan in the legacy source (not a real loan; skipped, harmless)',
}
BLOCKING = {'SCHEDULE_NOT_INFERABLE', 'MIXED_VOUCHER_INTEREST', 'CREDIT_NO_CASE', 'CREDIT_BEFORE_DISBURSED', 'CONSOLIDATION_UNRESOLVED', 'OTHER', 'LEDGER_OWES_NO_CASE', 'MANUAL_CORRECTION'}

def categorize(msg):
    if 'invalid loan_amt' in msg: return 'PLACEHOLDER_LOAN'
    if 'SCHEDULE_VERSION_NOT_BUILT' in msg: return 'SCHEDULE_NOT_INFERABLE'
    if msg.startswith('Mixed RLN/ALN voucher'): return 'MIXED_VOUCHER_INTEREST'
    if msg.startswith('MANUAL_CORRECTION'): return 'MANUAL_CORRECTION'
    if msg.startswith('Overpayment'): return 'OVERPAYMENT'
    if 'reconciled to ledger' in msg: return 'RECONCILED_TO_LEDGER'
    if msg.startswith('INFO'): return 'INFO'
    if 'LEDGER_OWES' in msg or 'LEDGER_BELOW' in msg: return 'LEDGER_OWES_NO_CASE'
    if 'Unattributed CR' in msg and 'not disbursed until' in msg: return 'CREDIT_BEFORE_DISBURSED'
    if 'Unattributed CR' in msg: return 'CREDIT_NO_CASE'
    if msg.startswith('Consolidation') or 'not found in Postgres' in msg: return 'CONSOLIDATION_UNRESOLVED'
    return 'OTHER'

def parse_log(path):
    """flags per member (from the SUMMARY section) and the set of members processed (member blocks)."""
    flags = defaultdict(list)
    processed = set()
    if not os.path.exists(path):
        return flags, processed
    text = open(path, encoding='utf8', errors='replace').read()
    processed = set(re.findall(r'^===== (\d+) =====', text, flags=re.M))
    i = text.find('--- Flags')
    for line in (text[i:] if i >= 0 else '').splitlines():
        m = re.match(r'^\[(\d+)\]\s+(.*)$', line.strip())
        if m:
            flags[m.group(1)].append((categorize(m.group(2)), m.group(2)))
    return flags, processed

FLAGS = defaultdict(list)
PROCESSED = set()
LIVE_LOGS = {'reset-pilot30-phase2.log', 'reset-demand-only-phase2.log', 'reset-others-clean-phase2.log'}
for fn in ['dryrun-other-members-phase2.log', 'reset-pilot30-phase2.log', 'reset-demand-only-phase2.log', 'reset-others-clean-phase2.log']:
    f, p = parse_log(os.path.join(R, fn))
    PROCESSED |= p
    for mb, v in f.items():
        FLAGS[mb] = v if fn not in LIVE_LOGS or mb not in FLAGS else v   # live result replaces an earlier dry-run result

after_idx = {(r['member_no'], r['loan_type']): r for r in after}

def cause_of(r):
    """Why an ACTIVE account's loan record differs from the ledger (the ledger is the legacy truth)."""
    led = float(r['ledger_balance']); rec = float(r['record_balance']); lag = float(r.get('lag_principal') or 0)
    if abs(led - rec) <= 10: return 'Matches'
    if lag > 0 and abs(led - (rec - lag)) <= 10: return 'Old-EMI kept aside in record (by design)'
    if rec == 0: return 'Record closed / invalid legacy loan row, ledger owes'
    if not r['has_schedule']: return 'No schedule built (data exception or not migrated)'
    if int(r['topups']) > 0: return 'Top-up / consolidation amounts differ'
    if int(r['journals']) > 0: return 'Journal entries in the ledger'
    return 'Other (schedule present)'

# Updated by hand as each pattern is checked (see backend/scripts/legacy-migration/migration.md, "Pattern verification tracker").
PATTERN_STATUS = {}
try:
    import json as _j
    PATTERN_STATUS = _j.load(open(os.path.join(HERE, 'pattern_status.json'), encoding='utf8'))
except Exception:
    PATTERN_STATUS = {}

ISSUE_KEYS = [k for _, k, _, _ in ISS]

def pats_of(r): return [pid for pid, k, _, _ in PATS if r.get(k)]
def issues_of(r): return [iid for iid, k, _, _ in ISS if r.get(k)]
def blocking_flags(mb): return [(c, m) for c, m in FLAGS.get(mb, []) if c in BLOCKING]
ISSUE_TEXT = {i: n for i, _, n, _ in ISS}

ST_FIXED = 'Fixed by re-migration'
ST_REMIG_MANUAL = 'Re-migrated - manual work needed'
ST_REMIG_CHECK = 'Re-migrated - difference remains, please check'
ST_MAN_EXC = 'Manual - data exception, not migrated'
ST_MAN_NONE = 'Manual - nothing to replay, not migrated'
ST_MAN_FAIL = 'Manual - earlier replay failed, not migrated'
ST_OK = 'No issue found'
ST_TODO = 'Not migrated yet'
STATUS_FILL = {ST_FIXED: 'C6EFCE', ST_OK: 'E2EFDA', ST_REMIG_CHECK: 'FFEB9C', ST_REMIG_MANUAL: 'F8CBAD', ST_MAN_EXC: 'FFC7CE', ST_MAN_NONE: 'FFC7CE', ST_MAN_FAIL: 'FFC7CE', ST_TODO: 'FFEB9C'}

def classify(b):
    mb, lt = b['member_no'], b['loan_type']
    reset = mb in reset_all
    a = after_idx.get((mb, lt)) if reset else None
    cur = a if a is not None else b
    open_issues = issues_of(cur)
    flags = blocking_flags(mb)
    reasons, details = [], []
    if flags:
        cats = sorted({c for c, _ in flags})
        reasons += [FLAG_TEXT[c] for c in cats]
        seen = Counter()
        for c, m in flags:
            if seen[c] < 2:
                details.append(m[:170]); seen[c] += 1
    if reset:
        left = [i for i in open_issues if i not in ('I1', 'I2')]
        if reset and a is None:
            status = ST_TODO
        elif not left and not flags:
            status = ST_FIXED
        elif flags:
            status = ST_REMIG_MANUAL
        else:
            status = ST_REMIG_CHECK
        reasons += [ISSUE_TEXT[i] for i in left]
    else:
        if 'I4' in open_issues:
            status = ST_MAN_FAIL; reasons.append('Earlier Phase 2 replay failed for this member (source-data exception)')
        elif flags:
            status = ST_MAN_EXC
        elif mb not in PROCESSED:
            status = ST_MAN_NONE; reasons.append('The legacy ledger holds no repayment receipts to replay for this member; the loan data cannot be rebuilt by the migration')
        elif any(i in open_issues for i in ('I1', 'I2', 'I3', 'I5', 'I6', 'I7')):
            status = ST_TODO
        else:
            status = ST_OK
    biz = []
    if cur.get('p07_silent_12_months'): biz.append('No payment for 12+ months: the customer must confirm whether the loan is unpaid or was paid elsewhere')
    if cur.get('p10_never_paid'): biz.append('Never paid since taken: the customer must confirm')
    if cur.get('p06_cash_lump_sum'): biz.append('Cash lump sum: business decision on shortening the term or lowering the EMI')
    nxt = {
        ST_FIXED: 'None', ST_OK: 'None',
        ST_REMIG_MANUAL: 'Settle the data exception by hand against the legacy ledger',
        ST_REMIG_CHECK: 'Compare ledger and loan record (Loan Position panel); correct with Journal/Transfer Entry or a repayment if needed',
        ST_MAN_EXC: 'Customer / operator reconciles the source data by hand; migrate afterwards',
        ST_MAN_NONE: 'Customer supplies or confirms the loan and payment history; then load it',
        ST_MAN_FAIL: 'Reconcile by hand against the legacy ledger; migrate afterwards',
        ST_TODO: 'Delete the member and re-migrate (delete / Phase 1 / Phase 2)',
    }[status]
    return status, nxt, reasons, details, biz, open_issues, cur

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(bold=True, color='FFFFFF')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

def sheet(wb, title, header, rows, widths=None, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=(c.column in wrap_cols), vertical='top')
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(40, max(12, len(str(h)) + 2)))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws

WORK_LOG = [
 ['2026-10-03', 'Fix', 'Replay script: first EMI due month = disbursement month + slot delay (was one month later); replayed payments tagged DEMAND / CASH; a re-run un-flags first EMIs wrongly marked as the old loan\'s deduction', 'Script', 'bulk_ledger_replay.ts, loan-repayment.service.ts', 'See migration.md'],
 ['2026-10-03', 'Re-migration', 'Member 610032487: delete, Phase 1, Phase 2', '1 member', 'Clean', 'backups/pre-single-member-reset-610032487-2026-10-03T18-09-32Z.dump'],
 ['2026-10-03', 'Re-migration', 'Member 610032274: delete, Phase 1, Phase 2', '1 member', 'Clean', 'backups/pre-single-member-reset-610032274-2026-10-03T18-29-18Z.dump'],
 ['2026-10-03', 'Re-migration', 'Pilot: 30 mixed members (15 consolidated, 15 not)', '30 members', '1,346 payments and 62 consolidations replayed, 210 flags, no errors', 'backups/pre-first-due-reset-pilot30-2026-10-03T18-35-38Z.dump'],
 ['2026-10-03', 'Re-migration', 'All demand-only pattern members that carried an issue (members whose earlier replay failed left untouched)', '1,003 members', '56,082 payments replayed, 5,685 flags; first-EMI-month fault removed from every account', 'backups/pre-first-due-reset-demand-only-2026-10-03T18-43-50Z.dump'],
 ['2026-10-04', 'Classification', 'Read-only replay run over the other members to find which ones have data exceptions (nothing written)', '1,775 members', '1,263 processed, 512 have no repayment receipts to replay; 1,033 with blocking flags', 'reports/dryrun-other-members-phase2.log'],
 ['2026-10-04', 'Re-migration', 'All other members whose replay raised no blocking flag', '227 members', 'See Read me sheet', 'see reports/reset-others-clean-status.log'],
]

def main():
    wb = Workbook()
    wb.remove(wb.active)
    accounts = before
    members = {r['member_no'] for r in accounts}
    classified = [(b,) + classify(b) for b in accounts]   # (b, status, nxt, reasons, details, biz, open_issues, cur)
    st_count = Counter(c[1] for c in classified)
    mem_status = {}
    for c in classified:
        mem_status.setdefault(c[0]['member_no'], set()).add(c[1])

    # --- Read me
    ws = wb.create_sheet('Read me')
    manual_accts = sum(v for k, v in st_count.items() if k.startswith('Manual') or k == ST_REMIG_MANUAL)
    lines = [
        ('Migration Pattern Register', True),
        ('One workbook for all legacy-migration work: patterns found, issues found, what was migrated, and what a person still has to do.', False),
        (f'Covers {len(accounts)} open loan accounts (member + loan type) over {len(members)} members. An account is open if the ledger or the loan record still shows money owed. After-scan used: {AFTER_LABEL}.', False),
        ('', False),
        ('Sheets', True),
        ('Work log - every migration step performed, with its scope, result and backup.', False),
        ('Patterns - 12 repayment patterns (P01-P12) and 7 migration issues (I1-I7): counts, what the migration can do, evidence from the re-migrated members.', False),
        ('Members - every account: patterns, issues, status, next action, reason.', False),
        ('Manual worklist - accounts that need a person, with the exact reason.', False),
        ('Manual reasons - the reasons grouped, with counts and what to do about each.', False),
        ('Replay flags - the data exceptions raised by the replay.', False),
        ('Demand-only result - before / after for the demand-only members.', False),
        ('Decisions - answers on one script vs one per pattern.', False),
        ('', False),
        ('Status values', True),
        (f'{ST_FIXED}: deleted from Postgres and migrated again from the legacy source; ledger and loan record agree.', False),
        (f'{ST_REMIG_CHECK}: re-migrated with no data exception, but ledger and loan record still differ (or there is no schedule) - look at the Loan Position panel.', False),
        (f'{ST_REMIG_MANUAL}: re-migrated, and the replay raised a data exception that only a person can settle.', False),
        (f'{ST_MAN_EXC}: the replay (trial) raised a data exception; left exactly as it was until the exception is settled.', False),
        (f'{ST_MAN_NONE}: the legacy ledger has no repayment receipts for this member, so nothing can be rebuilt.', False),
        (f'{ST_MAN_FAIL}: the earlier replay failed on a source-data exception.', False),
        ('', False),
        ('Status counts (accounts)', True),
    ] + [(f'  {k}: {v}', False) for k, v in st_count.most_common()] + [
        ('', False),
        (f'Members needing manual work: {sum(1 for m, s in mem_status.items() if any(x.startswith("Manual") or x == ST_REMIG_MANUAL for x in s))} of {len(members)}', True),
        ('', False),
        ('Refresh: run scripts/legacy-migration/_pattern_scan.js <label>, then python _build_pattern_register.py <label>. Procedures and incidents: backend/scripts/legacy-migration/migration.md', False),
    ]
    for text, bold in lines:
        ws.append([text])
        ws.cell(ws.max_row, 1).font = Font(bold=bold, size=14 if text == 'Migration Pattern Register' else 11)
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 150

    # --- Work log
    sheet(wb, 'Work log', ['Date', 'Kind', 'What was done', 'Scope', 'Result', 'Backup / log'], WORK_LOG,
          widths={1: 12, 2: 14, 3: 90, 4: 14, 5: 60, 6: 70}, wrap_cols=(3, 5, 6))

    # --- Patterns
    def cnt(key, lt=None): return sum(1 for r in accounts if r[key] and (lt is None or r['loan_type'] == lt))
    def mem(key): return len({r['member_no'] for r in accounts if r[key]})
    def evidence(key):
        rows = [c for c in classified if c[0][key] and c[0]['member_no'] in reset_all and (c[0]['member_no'], c[0]['loan_type']) in after_idx]
        fixed = sum(1 for c in rows if c[1] == ST_FIXED)
        manual = sum(1 for c in rows if c[1] == ST_REMIG_MANUAL)
        return len(rows), fixed, manual
    HAND = {
        'P01': ('Yes', 'Demand money is tagged DEMAND and bucketed to its own month.'),
        'P02': ('Partly', 'Cash receipts are replayed and pooled oldest-first; they are not tied to a month.'),
        'P03': ('Yes', 'Demand pays its own month; cash on top is pooled.'),
        'P04': ('Partly', 'Top-ups are detected from matching journal pairs; ambiguous ones are flagged for hand review.'),
        'P05': ('Yes', 'A new schedule version is built at each top-up; the EMI seed is the first normal recurring principal.'),
        'P06': ('Needs decision', 'The money is replayed; whether the term shortens or the EMI falls is a business rule.'),
        'P07': ('Needs decision', 'The migration copies what the source holds; only the customer knows whether the loan is unpaid or paid elsewhere.'),
        'P08': ('Partly', 'Journals are copied to the ledger and shown in the Loan Position panel, but they do not change the loan record.'),
        'P09': ('Partly', 'Consolidations are detected from journal pairs and replayed; unmatched ones are flagged.'),
        'P10': ('Needs decision', 'Nothing to replay: the customer confirms whether the loan was ever paid.'),
        'P11': ('Yes', 'Missed months stay as arrears in their own months and earn penalty from the real due month.'),
        'P12': ('Partly', 'The replay copies the ledger; interest that was never collected cannot be invented.'),
    }
    HAND_I = {
        'I1': ('Yes', 'Fixed in the replay script; delete + Phase 1 + Phase 2 corrects the stored date. Removed from every re-migrated account.'),
        'I2': ('Yes', "Disappears with I1 (the first EMI is no longer treated as the old loan's deduction)."),
        'I3': ('Partly', 'A re-run builds a schedule when a normal recurring principal exists; otherwise the replay flags it (manual).'),
        'I4': ('No', 'Source-data exception; a person must reconcile the loan against the legacy ledger.'),
        'I5': ('Partly', 'Closes when the payments are in the ledger; otherwise manual.'),
        'I6': ('Partly', 'Differences caused by I1/I2 close; those from old-EMI deductions, journals, cash or unattributed payments remain.'),
        'I7': ('No', 'The loan record is closed but the ledger still owes: reconcile by hand.'),
    }
    rows = []
    for pid, key, name, desc in PATS:
        n, fixed, manual = evidence(key)
        h, why = HAND[pid]
        rows.append([pid, 'Repayment pattern', name, desc, cnt(key, 'ALN'), cnt(key, 'RLN'), cnt(key), mem(key), h, why, n, fixed,
                     round(100.0 * fixed / n) if n else None, manual,
                     'Customer decision' if h == 'Needs decision' else 'One script is enough; flagged accounts go to the manual list'])
    for iid, key, name, desc in ISS:
        n, fixed, manual = evidence(key)
        h, why = HAND_I[iid]
        rows.append([iid, 'Migration issue', name, desc, cnt(key, 'ALN'), cnt(key, 'RLN'), cnt(key), mem(key), h, why, n, fixed,
                     round(100.0 * fixed / n) if n else None, manual,
                     'Manual' if h == 'No' else 'One script is enough; flagged accounts go to the manual list'])
    sheet(wb, 'Patterns', ['ID', 'Kind', 'Name', 'How it is recognised', 'Accounts ALN', 'Accounts RLN', 'Accounts total', 'Members', 'Can the migration fix it?', 'Why',
                           'Re-migrated accounts (evidence)', 'Of those: fully clean', 'Clean %', 'Of those: manual work needed', 'Dedicated script / customer action'],
          rows, widths={1: 6, 2: 17, 3: 34, 4: 58, 5: 10, 6: 10, 7: 10, 8: 9, 9: 15, 10: 66, 11: 13, 12: 11, 13: 8, 14: 13, 15: 40}, wrap_cols=(3, 4, 10, 15))

    # --- Flags
    fl_members = defaultdict(set)
    for mb, items in FLAGS.items():
        for c, _ in items:
            fl_members[c].add(mb)
    fl_rows = [[c, FLAG_TEXT[c], len(fl_members.get(c, ())), 'Blocking - manual' if c in BLOCKING else 'Information only',
                ', '.join(sorted(fl_members.get(c, ()))[:40]) + (' ...' if len(fl_members.get(c, ())) > 40 else '')] for c in FLAG_TEXT]
    sheet(wb, 'Replay flags', ['Flag type', 'Meaning', 'Members', 'Effect', 'Sample member numbers'], fl_rows, widths={1: 28, 2: 90, 3: 10, 4: 20, 5: 80}, wrap_cols=(2, 5))

    # --- Members
    mrows = []
    for b, st, nxt, reasons, details, biz, op, cur in classified:
        reset = b['member_no'] in reset_all
        mrows.append([b['member_no'], b['member_name'], b['loan_type'], cur.get('loan_case') or b.get('loan_case'), b.get('loan_amt'), b.get('no_of_instal'), b.get('instal_amt'), b.get('case_date'),
                      float(cur['ledger_balance']), float(cur['record_balance']), float(cur['difference']), cur.get('last_payment'),
                      cur['demand_payments'], cur['cash_payments'], cur['topups'], cur['journals'],
                      ', '.join(pats_of(b)), ', '.join(issues_of(b)), (', '.join(op) if reset else ''),
                      group_of.get(b['member_no'], 'No'), st, nxt, '; '.join(reasons), ' | '.join(details), '; '.join(biz)])
    ws = sheet(wb, 'Members', ['Member no', 'Member name', 'Loan type', 'Loan case', 'Loan amount', 'EMIs', 'EMI (stored)', 'Case date', 'Ledger balance (now)', 'Loan record balance (now)', 'Difference (now)',
                               'Last payment', 'Demand payments', 'Cash payments', 'Top-ups', 'Journals', 'Patterns', 'Issues before', 'Issues left after re-migration', 'Re-migrated in',
                               'Status', 'Next action', 'Reason (why manual / what remains)', 'Example data exceptions', 'Customer / business decision notes'],
               mrows, widths={1: 12, 2: 26, 3: 8, 4: 9, 5: 12, 6: 6, 7: 10, 8: 11, 9: 13, 10: 13, 11: 12, 12: 12, 13: 9, 14: 9, 15: 7, 16: 8, 17: 24, 18: 14, 19: 16, 20: 30, 21: 36, 22: 46, 23: 80, 24: 80, 25: 60},
               wrap_cols=(21, 22, 23, 24, 25))
    for row in ws.iter_rows(min_row=2):
        c = row[20]
        if c.value in STATUS_FILL:
            c.fill = PatternFill('solid', fgColor=STATUS_FILL[c.value])

    # --- Manual worklist (+ reasons)
    man = [c for c in classified if c[1].startswith('Manual') or c[1] in (ST_REMIG_MANUAL, ST_REMIG_CHECK)]
    sheet(wb, 'Manual worklist', ['Member no', 'Member name', 'Loan type', 'Loan case', 'Ledger balance', 'Loan record balance', 'Difference', 'Patterns', 'Status', 'Reason', 'Example data exceptions', 'Customer / business decision notes'],
          [[c[0]['member_no'], c[0]['member_name'], c[0]['loan_type'], c[7].get('loan_case') or c[0].get('loan_case'), float(c[7]['ledger_balance']), float(c[7]['record_balance']), float(c[7]['difference']),
            ', '.join(pats_of(c[0])), c[1], '; '.join(c[3]) or c[2], ' | '.join(c[4]), '; '.join(c[5])] for c in man],
          widths={1: 12, 2: 26, 3: 8, 4: 9, 5: 13, 6: 13, 7: 12, 8: 24, 9: 38, 10: 90, 11: 80, 12: 60}, wrap_cols=(9, 10, 11, 12))
    reason_count = Counter()
    reason_members = defaultdict(set)
    for c in man:
        for r in (c[3] or [c[2]]):
            reason_count[r] += 1; reason_members[r].add(c[0]['member_no'])
    sheet(wb, 'Manual reasons', ['Reason', 'Accounts', 'Members'], [[r, n, len(reason_members[r])] for r, n in reason_count.most_common()], widths={1: 130, 2: 10, 3: 10}, wrap_cols=(1,))

    # --- Demand-only result
    d_rows = [c for c in classified if c[0]['p01_demand_only']]
    res = Counter(c[1] for c in d_rows)
    sheet(wb, 'Demand-only result', ['Measure', 'Accounts'],
          [['Demand-only accounts', len(d_rows)], ['Members', len({c[0]['member_no'] for c in d_rows})]] +
          [[f'{iid} before ({n})', sum(1 for c in d_rows if c[0][k])] for iid, k, n, _ in ISS] +
          [[f'{iid} after re-migration ({n})', sum(1 for c in d_rows if c[0]['member_no'] in reset_all and (c[0]['member_no'], c[0]['loan_type']) in after_idx and after_idx[(c[0]['member_no'], c[0]['loan_type'])][k])] for iid, k, n, _ in ISS] +
          [[f'Status: {k}', v] for k, v in res.most_common()], widths={1: 80, 2: 12})

    # --- Decisions
    sheet(wb, 'Decisions', ['Question', 'Answer', 'Evidence'], [
        ['Is one migration script enough, or does each pattern need its own?',
         'One script. The faults found (first EMI month one month late, first EMI marked as the old loan\'s deduction) sit in the shared replay logic, not in any pattern. Patterns differ only in the data exceptions they produce, and those go to the manual list.',
         'After re-migrating the demand-only members, I1 and I2 remain on 0 accounts; see the Patterns sheet for the other patterns.'],
        ['Which things cannot be handled by migration?',
         'Data exceptions flagged by the replay (no loan to apply a payment to, payment before disbursement, mixed regular/emergency voucher, schedule not inferable, unmatched consolidation), members with no repayment receipts, and customer decisions (silent loans, lump sums, never paid).',
         'Manual reasons sheet.'],
        ['What do we tell the customer?',
         'For the Manual worklist: the legacy data for these loans is not consistent (see the Reason column); they must confirm the loan terms and payment history, or we correct them by hand later. For silent / never-paid / lump-sum loans: confirm whether the loan is really unpaid and how lump sums should shorten the loan.',
         'Manual worklist sheet.'],
        ['Is anything else needed besides the migration script?',
         'A post-migration reconciliation view (the Loan Position panel already shows ledger vs loan record) and, later, a rule for loans whose schedule cannot be inferred (customer supplies the EMI). No per-pattern migration script is needed.',
         'Loan Position panel; Replay flags sheet.'],
    ], widths={1: 52, 2: 120, 3: 70}, wrap_cols=(1, 2, 3))

    # --- One sheet per repayment pattern: ACTIVE loans only (current DB state = after scan), member ID + loan case + manual flag
    act = {}
    for r in after:
        if float(r['ledger_balance']) > 100:   # active = the ledger still shows money owed
            k = (r['member_no'], r['loan_type'])
            if k not in act or (not act[k].get('loan_case') and r.get('loan_case')):
                act[k] = r
    cl = {k: classify(r) for k, r in act.items()}
    # Loans the loan record still lists as open while the ledger says repaid (not active; shown in the closure window until fixed)
    ro = sorted({(r['member_no'], r['loan_type'], r.get('loan_case') or '', float(r['record_balance'])) for r in after
                 if float(r['record_balance']) > 0 and float(r['ledger_balance']) <= 100}, key=lambda x: (int(x[0]), x[1]))
    sheet(wb, 'Repaid - record still open', ['Member ID', 'Loan type', 'Loan case no', 'Record balance (wrong)', 'Re-migrated'],
          [[m, t, c, b, 'Yes' if m in reset_all else 'No'] for m, t, c, b in ro], widths={1: 14, 2: 10, 3: 14, 4: 20, 5: 12})
    MAN_YES = {ST_MAN_EXC, ST_MAN_NONE, ST_MAN_FAIL, ST_REMIG_MANUAL}
    def manual_of(st):
        if st in MAN_YES: return 'YES - manual'
        if st == ST_REMIG_CHECK: return 'CHECK - difference remains'
        if st == ST_TODO: return 'PENDING - not migrated yet'
        return 'No'
    FILL = {'YES - manual': 'FFC7CE', 'CHECK - difference remains': 'FFEB9C', 'PENDING - not migrated yet': 'DDEBF7'}
    summ = []
    for pid, key, name, _ in PATS:
        recs = sorted((k for k, r in act.items() if r[key]), key=lambda k: (int(k[0]), k[1]))
        rows = []
        for k in recs:
            st, nxt, reasons, details, biz, op, cur = cl[k]
            rows.append([k[0], k[1], act[k].get('loan_case') or '', manual_of(st), cause_of(act[k]), '; '.join(reasons)])
        ws = sheet(wb, f'{pid} {name}'[:31].replace('/', '-'), ['Member ID', 'Loan type', 'Active loan case no', 'Manual work?', 'Record vs ledger', 'Reason (why manual)'], rows,
                   widths={1: 14, 2: 10, 3: 18, 4: 28, 5: 44, 6: 110}, wrap_cols=(5, 6))
        for row in ws.iter_rows(min_row=2):
            f = FILL.get(row[3].value)
            if f:
                for c in row: c.fill = PatternFill('solid', fgColor=f)
        summ.append([pid, name, len(rows), sum(1 for r in rows if r[3].startswith('YES')), sum(1 for r in rows if r[3].startswith('CHECK')), sum(1 for r in rows if r[3].startswith('PENDING'))])
    CAUSES = ['Matches', 'Old-EMI kept aside in record (by design)', 'Record closed / invalid legacy loan row, ledger owes',
              'No schedule built (data exception or not migrated)', 'Top-up / consolidation amounts differ', 'Journal entries in the ledger', 'Other (schedule present)']
    trows = []
    for pid, key, name, _ in PATS:
        recs = [r for r in act.values() if r[key]]
        cc = Counter(cause_of(r) for r in recs)
        st = PATTERN_STATUS.get(pid, {})
        trows.append([pid, name, len(recs)] + [cc.get(c, 0) for c in CAUSES] + [round(100.0 * cc.get('Matches', 0) / len(recs)) if recs else None,
                      st.get('status', 'Not checked yet'), st.get('can_migration_fix', ''), st.get('finding', ''), st.get('checked_on', '')])
    ws = sheet(wb, 'Pattern tracker', ['ID', 'Pattern', 'Active loans'] + CAUSES + ['Matches %', 'Check status', 'Can migration fix it?', 'Finding / what was done', 'Last checked'],
               trows, widths={1: 6, 2: 34, 3: 10, 4: 10, 5: 16, 6: 18, 7: 18, 8: 14, 9: 14, 10: 14, 11: 10, 12: 26, 13: 22, 14: 90, 15: 12}, wrap_cols=(2, 14))
    wb.move_sheet('Pattern tracker', offset=-(len(wb.sheetnames) - 2))
    print('active loans', len(act), 'with any pattern', len({k for k, r in act.items() if any(r[x] for _, x, _, _ in PATS)}))
    for r in summ: print(r)

    # --- Sample 50: the fixed test sample used for every migration round (see migration.md)
    try:
        SL = os.environ.get('SAMPLE_LABEL', 'sample-r11')
        sample = [l.strip() for l in open(os.path.join(R, 'sample50-members.txt')) if l.strip()]
        scan = {(r['member_no'], r['loan_type']): r for r in json.load(open(os.path.join(R, f'pattern-scan-{SL}.json')))}
        audit = {(a['member'], a['head']): a for a in json.load(open(os.path.join(R, f'ledger-audit-{SL}.json')))}
        flagtxt = defaultdict(list)
        lg = open(os.path.join(R, f'reset-sample50-{SL.split("-")[-1]}-phase2.log'), encoding='utf8').read()
        i0 = lg.find('--- Flags')
        for line in (lg[i0:] if i0 >= 0 else '').splitlines():
            m = re.match(r'^\[(\d+)\]\s+(.*)$', line.strip())
            if m: flagtxt[m.group(1)].append(categorize(m.group(2)))
        names = {}
        for r in scan.values(): names[r['member_no']] = r.get('member_name') or ''
        fin = {}
        try:
            for x in json.load(open(os.path.join(R, 'sample-final-sample-final.json'), encoding='utf8')): fin[x['member']] = x
        except Exception: pass
        srows = []
        for mb in sample:
            keys = [k for k in scan if k[0] == mb] or [(mb, '')]
            for k in keys:
                r = scan.get(k, {})
                a = audit.get(k, {})
                active = float(r.get('ledger_balance', 0) or 0) > 100 if r else False
                pats = [pid for pid, key, _, _ in PATS if r.get(key)]
                fl = Counter(flagtxt.get(mb, []))
                srows.append([mb, names.get(mb, ''), k[1], r.get('loan_case') or '', 'Active' if active else 'Not active',
                              ', '.join(pats), float(r.get('ledger_balance', 0) or 0), float(r.get('record_balance', 0) or 0),
                              'Yes' if a.get('closingMatches') else ('n/a' if not a else 'No'),
                              len(a.get('unjustifiedCredits', [])) if a else '',
                              'Yes' if r.get('has_schedule') else 'No',
                              '; '.join(f'{c} x{n}' for c, n in fl.most_common() if c not in ('PLACEHOLDER_LOAN',)),
                              fin.get(mb, {}).get('verdict', ''), ' | '.join(fin.get(mb, {}).get('reasons', []))])
        ws = sheet(wb, 'Sample 50', ['Member ID', 'Member name', 'Loan type', 'Loan case', 'Status', 'Patterns', 'Ledger balance', 'Record balance',
                                     'Closing balance matches', 'Unjustified credits', 'Schedule built', 'Replay flags (type x members count)', 'Verdict (final check)', 'Reasons'],
                   srows, widths={1: 12, 2: 26, 3: 9, 4: 10, 5: 11, 6: 30, 7: 14, 8: 14, 9: 12, 10: 11, 11: 10, 12: 60, 13: 18, 14: 90}, wrap_cols=(6, 12, 14))
        for row in ws.iter_rows(min_row=2):
            v = row[12].value
            col = {'CLEAN': 'C6EFCE', 'TO CONFIRM': 'FFEB9C', 'FLAGGED - MANUAL': 'FFC7CE'}.get(v)
            if col: row[12].fill = PatternFill('solid', fgColor=col)
            if row[8].value == 'No':
                for c in row: c.fill = PatternFill('solid', fgColor='FFC7CE')
        print('sample rows', len(srows))
    except Exception as ex:
        print('sample sheet skipped:', ex)

    # --- Manual corrections register (re-applied on every re-migration)
    try:
        corr = json.load(open(os.path.join(HERE, 'manual-corrections.json'), encoding='utf8'))
        head = {'A1002': 'Regular loan (RLN)', 'A1047': 'Emergency loan (ALN)'}
        sheet(wb, 'Manual corrections', ['Member ID', 'Voucher', 'Date', 'Credit posted to', 'Belongs to', 'Amount', 'Status', 'Reason'],
              [[c['mbno'], c['voucher'], c['date'], head.get(c['code'], c['code']), head.get(c['toCode'], c['toCode']), c['amount'],
                'Applied by the migration (re-applied on every re-migration); flagged for manual review', c['reason']] for c in corr],
              widths={1: 12, 2: 10, 3: 12, 4: 22, 5: 22, 6: 10, 7: 50, 8: 120}, wrap_cols=(7, 8))
    except Exception as ex:
        print('corrections sheet skipped:', ex)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print('saved', OUT)
    print(dict(st_count))
    print('members needing manual work:', sum(1 for m, s in mem_status.items() if any(x.startswith('Manual') or x == ST_REMIG_MANUAL for x in s)), 'of', len(members))

main()
